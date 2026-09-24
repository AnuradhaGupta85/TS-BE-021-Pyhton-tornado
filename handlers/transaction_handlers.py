# Transaction CRUD, filtering, sorting, and monthly summary handlers.
import csv
import io
import os
import uuid
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from dotenv import load_dotenv
from pydantic import ValidationError
from sqlalchemy import func, or_, select
from sqlalchemy.orm import selectinload
import tornado.web

from base_handler import BaseHandler
from models import Category, Tag, Transaction, TransactionAttachment, TransactionType
from schemas import TransactionCreateSchema, TransactionResponse, TransactionUpdateSchema

load_dotenv('.env_10c6229a-2c2d-44bf-a722-7f4465657ae8', override=True)


def transaction_data(row: Transaction) -> dict:
    return TransactionResponse(id=row.id, amount=row.amount, type=row.type, category_id=row.category_id, category_name=row.category.name, date=row.date, description=row.description, user_id=row.user_id, tags=[{'id': tag.id, 'name': tag.name} for tag in row.tags], attachments=[{'id': attachment.id, 'filename': attachment.filename, 'content_type': attachment.content_type} for attachment in row.attachments]).model_dump(mode='json')


async def accessible_tags(handler: BaseHandler, tag_ids: list[int]) -> list[Tag]:
    if not tag_ids:
        return []
    if len(tag_ids) != len(set(tag_ids)):
        raise tornado.web.HTTPError(422, reason='tag_ids must not contain duplicates')
    tags = (await handler.db.execute(select(Tag).where(Tag.id.in_(tag_ids), Tag.owner_id == handler.current_user.id))).scalars().all()
    if len(tags) != len(tag_ids):
        raise tornado.web.HTTPError(422, reason='One or more tags are unavailable')
    return tags


async def accessible_category(handler: BaseHandler, category_id: int) -> Category:
    category = (await handler.db.execute(select(Category).where(Category.id == category_id))).scalar_one_or_none()
    if category is None or (not category.is_predefined and category.owner_id != handler.current_user.id):
        raise tornado.web.HTTPError(422, reason='Category is unavailable')
    return category


class TransactionCollectionHandler(BaseHandler):
    requires_auth = True

    async def get(self) -> None:
        limit, offset = self.pagination()
        conditions = [Transaction.user_id == self.current_user.id]
        search = self.get_query_argument('search', '').strip()
        if search:
            conditions.append(or_(Transaction.description.ilike(f'%{search}%'), Category.name.ilike(f'%{search}%')))
        type_value = self.get_query_argument('type', '').strip()
        if type_value:
            try:
                conditions.append(Transaction.type == TransactionType(type_value))
            except ValueError as exc:
                raise tornado.web.HTTPError(400, reason='type must be Income or Expense') from exc
        for param, column, label in [('start_date', Transaction.date, 'start_date'), ('end_date', Transaction.date, 'end_date')]:
            raw = self.get_query_argument(param, '').strip()
            if raw:
                try:
                    parsed = date.fromisoformat(raw)
                except ValueError as exc:
                    raise tornado.web.HTTPError(400, reason=f'{label} must be YYYY-MM-DD') from exc
                conditions.append(column >= parsed if param == 'start_date' else column <= parsed)
        category_id = self.get_query_argument('category_id', '').strip()
        if category_id:
            try:
                conditions.append(Transaction.category_id == int(category_id))
            except ValueError as exc:
                raise tornado.web.HTTPError(400, reason='category_id must be an integer') from exc
        tag_id = self.get_query_argument('tag_id', '').strip()
        if tag_id:
            try:
                conditions.append(Transaction.tags.any(Tag.id == int(tag_id)))
            except ValueError as exc:
                raise tornado.web.HTTPError(400, reason='tag_id must be an integer') from exc
        sort_by = self.get_query_argument('sort_by', 'date')
        sort_fields = {'date': Transaction.date, 'amount': Transaction.amount, 'category': Category.name}
        if sort_by not in sort_fields:
            raise tornado.web.HTTPError(400, reason='sort_by must be date, amount, or category')
        order = self.get_query_argument('order', 'desc').lower()
        if order not in {'asc', 'desc'}:
            raise tornado.web.HTTPError(400, reason='order must be asc or desc')
        total = (await self.db.execute(select(func.count()).select_from(Transaction).join(Category).where(*conditions))).scalar_one()
        ordering = sort_fields[sort_by].asc() if order == 'asc' else sort_fields[sort_by].desc()
        rows = (await self.db.execute(select(Transaction).join(Transaction.category).options(selectinload(Transaction.category), selectinload(Transaction.tags), selectinload(Transaction.attachments)).where(*conditions).order_by(ordering, Transaction.id.desc()).limit(limit).offset(offset))).scalars().all()
        self.write_json({'items': [transaction_data(row) for row in rows], 'total': total, 'limit': limit, 'offset': offset})

    async def post(self) -> None:
        try:
            payload = TransactionCreateSchema.model_validate(self.json_body())
        except ValidationError as exc:
            raise tornado.web.HTTPError(422, reason=self.validation_error_reason(exc)) from exc
        await accessible_category(self, payload.category_id)
        tags = await accessible_tags(self, payload.tag_ids)
        row = Transaction(amount=payload.amount, type=payload.type, category_id=payload.category_id, date=payload.date, description=payload.description, user_id=self.current_user.id, tags=tags)
        self.db.add(row)
        await self.db.commit()
        result = await self.db.execute(select(Transaction).options(selectinload(Transaction.category), selectinload(Transaction.tags), selectinload(Transaction.attachments)).where(Transaction.id == row.id))
        row = result.scalar_one()
        self.write_json(transaction_data(row), 201)


class TransactionDetailHandler(BaseHandler):
    requires_auth = True

    async def _row(self, transaction_id: str) -> Transaction:
        try:
            identifier = int(transaction_id)
        except ValueError as exc:
            raise tornado.web.HTTPError(404, reason='Transaction not found') from exc
        row = (await self.db.execute(select(Transaction).options(selectinload(Transaction.category), selectinload(Transaction.tags), selectinload(Transaction.attachments)).where(Transaction.id == identifier, Transaction.user_id == self.current_user.id))).scalar_one_or_none()
        if row is None:
            raise tornado.web.HTTPError(404, reason='Transaction not found')
        return row

    async def get(self, transaction_id: str) -> None:
        self.write_json(transaction_data(await self._row(transaction_id)))

    async def put(self, transaction_id: str) -> None:
        row = await self._row(transaction_id)
        try:
            payload = TransactionUpdateSchema.model_validate(self.json_body())
        except ValidationError as exc:
            raise tornado.web.HTTPError(422, reason=self.validation_error_reason(exc)) from exc
        values = payload.model_dump(exclude_unset=True)
        if 'category_id' in values:
            await accessible_category(self, values['category_id'])
        if 'tag_ids' in values:
            row.tags = await accessible_tags(self, values.pop('tag_ids') or [])
        for field, value in values.items():
            setattr(row, field, value)
        await self.db.commit()
        result = await self.db.execute(select(Transaction).options(selectinload(Transaction.category), selectinload(Transaction.tags), selectinload(Transaction.attachments)).where(Transaction.id == row.id))
        self.write_json(transaction_data(result.scalar_one()))

    async def delete(self, transaction_id: str) -> None:
        row = await self._row(transaction_id)
        await self.db.delete(row)
        await self.db.commit()
        self.set_status(204)
        self.finish()


class TransactionBulkHandler(BaseHandler):
    requires_auth = True

    async def post(self) -> None:
        files = self.request.files.get('file', [])
        if not files:
            raise tornado.web.HTTPError(400, reason='A CSV or Excel file is required in the file field')
        upload = files[0]
        if len(upload.body) > 10 * 1024 * 1024:
            raise tornado.web.HTTPError(413, reason='Upload must not exceed 10 MB')
        filename = upload.filename.lower()
        try:
            if filename.endswith('.csv'):
                records = list(csv.DictReader(io.StringIO(upload.body.decode('utf-8-sig'))))
            elif filename.endswith(('.xlsx', '.xlsm')):
                worksheet = load_workbook(io.BytesIO(upload.body), read_only=True, data_only=True).active
                headers = [str(value).strip() if value is not None else '' for value in next(worksheet.iter_rows(values_only=True))]
                records = [dict(zip(headers, row)) for row in worksheet.iter_rows(min_row=2, values_only=True) if any(value is not None for value in row)]
            else:
                raise tornado.web.HTTPError(422, reason='Only CSV and .xlsx files are supported')
        except tornado.web.HTTPError:
            raise
        except Exception as exc:
            raise tornado.web.HTTPError(422, reason='Unable to read the import file') from exc
        if not records:
            raise tornado.web.HTTPError(422, reason='The import file contains no transaction rows')
        if len(records) > 1000:
            raise tornado.web.HTTPError(422, reason='A bulk import may contain at most 1000 rows')
        rows = []
        for index, record in enumerate(records, start=2):
            try:
                payload = TransactionCreateSchema.model_validate({key: value for key, value in record.items() if value is not None})
                await accessible_category(self, payload.category_id)
                rows.append(Transaction(amount=payload.amount, type=payload.type, category_id=payload.category_id, date=payload.date, description=payload.description, user_id=self.current_user.id, tags=await accessible_tags(self, payload.tag_ids)))
            except (ValidationError, tornado.web.HTTPError) as exc:
                reason = self.validation_error_reason(exc) if isinstance(exc, ValidationError) else exc.reason
                raise tornado.web.HTTPError(422, reason=f'Row {index}: {reason}') from exc
        self.db.add_all(rows)
        await self.db.commit()
        self.write_json({'imported': len(rows)}, 201)


class TransactionAttachmentHandler(BaseHandler):
    requires_auth = True

    async def post(self, transaction_id: str) -> None:
        try:
            identifier = int(transaction_id)
        except ValueError as exc:
            raise tornado.web.HTTPError(404, reason='Transaction not found') from exc
        transaction = (await self.db.execute(select(Transaction).where(Transaction.id == identifier, Transaction.user_id == self.current_user.id))).scalar_one_or_none()
        if transaction is None:
            raise tornado.web.HTTPError(404, reason='Transaction not found')
        files = self.request.files.get('file', [])
        if not files:
            raise tornado.web.HTTPError(400, reason='An attachment is required in the file field')
        upload = files[0]
        if len(upload.body) > 10 * 1024 * 1024:
            raise tornado.web.HTTPError(413, reason='Attachment must not exceed 10 MB')
        extension = os.path.splitext(upload.filename)[1].lower()
        if extension not in {'.jpg', '.jpeg', '.png', '.pdf'}:
            raise tornado.web.HTTPError(422, reason='Only JPG, PNG, and PDF attachments are supported')
        uploads = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'uploads', 'receipts')
        os.makedirs(uploads, exist_ok=True)
        stored_filename = f'{uuid.uuid4().hex}{extension}'
        with open(os.path.join(uploads, stored_filename), 'wb') as output:
            output.write(upload.body)
        attachment = TransactionAttachment(transaction_id=transaction.id, filename=upload.filename, stored_filename=stored_filename, content_type=upload.content_type or 'application/octet-stream')
        self.db.add(attachment)
        await self.db.commit()
        self.write_json({'id': attachment.id, 'filename': attachment.filename, 'content_type': attachment.content_type}, 201)


class TagCollectionHandler(BaseHandler):
    requires_auth = True

    async def get(self) -> None:
        tags = (await self.db.execute(select(Tag).where(Tag.owner_id == self.current_user.id).order_by(Tag.name))).scalars().all()
        self.write_json({'items': [{'id': tag.id, 'name': tag.name} for tag in tags]})

    async def post(self) -> None:
        from schemas import TagCreateSchema
        try:
            payload = TagCreateSchema.model_validate(self.json_body())
        except ValidationError as exc:
            raise tornado.web.HTTPError(422, reason=self.validation_error_reason(exc)) from exc
        existing = (await self.db.execute(select(Tag).where(Tag.owner_id == self.current_user.id, Tag.name == payload.name))).scalar_one_or_none()
        if existing:
            raise tornado.web.HTTPError(409, reason='Tag already exists')
        tag = Tag(name=payload.name, owner_id=self.current_user.id)
        self.db.add(tag)
        await self.db.commit()
        self.write_json({'id': tag.id, 'name': tag.name}, 201)


class TagDetailHandler(BaseHandler):
    requires_auth = True

    async def delete(self, tag_id: str) -> None:
        try:
            identifier = int(tag_id)
        except ValueError as exc:
            raise tornado.web.HTTPError(404, reason='Tag not found') from exc
        tag = (await self.db.execute(select(Tag).where(Tag.id == identifier, Tag.owner_id == self.current_user.id))).scalar_one_or_none()
        if tag is None:
            raise tornado.web.HTTPError(404, reason='Tag not found')
        await self.db.delete(tag)
        await self.db.commit()
        self.set_status(204)
        self.finish()


class MonthlySummaryHandler(BaseHandler):
    requires_auth = True

    async def get(self) -> None:
        try:
            year = int(self.get_query_argument('year'))
            month = int(self.get_query_argument('month'))
            period_start = date(year, month, 1)
        except (tornado.web.MissingArgumentError, ValueError) as exc:
            raise tornado.web.HTTPError(400, reason='year and month must form a valid month') from exc
        period_end = date(year + (month == 12), 1 if month == 12 else month + 1, 1)
        base = [Transaction.user_id == self.current_user.id, Transaction.date >= period_start, Transaction.date < period_end]
        totals = (await self.db.execute(select(Transaction.type, func.coalesce(func.sum(Transaction.amount), 0)).where(*base).group_by(Transaction.type))).all()
        by_type = {kind: Decimal(str(value)) for kind, value in totals}
        breakdown = (await self.db.execute(select(Category.name, func.sum(Transaction.amount)).join(Transaction).where(*base, Transaction.type == TransactionType.EXPENSE).group_by(Category.id, Category.name).order_by(Category.name))).all()
        income = by_type.get(TransactionType.INCOME, Decimal('0'))
        expense = by_type.get(TransactionType.EXPENSE, Decimal('0'))
        self.write_json({'total_income': income, 'total_expenses': expense, 'balance': income - expense, 'expenses_by_category': [{'category': name, 'total': total} for name, total in breakdown]})
